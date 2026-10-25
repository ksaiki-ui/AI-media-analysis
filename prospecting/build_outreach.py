"""対象リストに会社情報・電話番号・フォームの調査結果を足し、架電用とフォーム用のシートにする。

使い方:
  python3 prospecting/build_outreach.py target_list.xlsx enriched.csv gbiz.csv jassa.csv out.xlsx

- target_list.xlsx: build_list.py の出力（区分・優先度・一言目・フォーム文面）
- enriched.csv: contact_enrich.py の出力（電話番号・フォームの中身）
- gbiz.csv: domain と gBizINFO の会社情報（正式社名・代表者・従業員数・所在地・法人番号・事業概要）
- jassa.csv: universe/jassa.csv（gBizINFO に無い会社の所在地）
"""

from __future__ import annotations

import csv
import sys
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TARGET = ("少量×更新中", "少量×停止", "大量×停止")
CALL_STATUS = [
    "未着手",
    "不在・再架電",
    "担当者不明",
    "断り",
    "資料送付",
    "アポ獲得",
    "対象外",
]
FORM_STATUS = ["未送信", "送信済", "送信できず", "送らない"]
SIGNATURE = "\n――――――――\n株式会社MIRAIKU\n【担当者名】\n【メール】\n【電話】\n"


def dom(u: str) -> str:
    return (
        urlparse(u if "://" in u else "https://" + u)
        .netloc.lower()
        .removeprefix("www.")
    )


def send_route(c: dict) -> tuple[str, str]:
    """(送り方, 理由)。自動で送る候補は、自社フォームで CAPTCHA なし・営業お断りなしのものだけ"""
    if c.get("no_sales_notice") == "True":
        return "送らない", "営業お断りの表記あり"
    kind = c.get("form_kind", "")
    if kind in ("見つからず", ""):
        return "フォームなし", "問い合わせフォームが見つからない（電話で当たる）"
    if kind == "メールのみ":
        return "メールのみ", "フォームが無くメールアドレスのみ"
    if c.get("captcha"):
        return "人が送信", f"{c['captcha']} あり"
    if kind.startswith("外部"):
        return "人が送信", f"外部のフォーム（{kind[3:-1]}）"
    if c.get("confirm_step") == "True":
        return "自動の候補（確認画面あり）", "自社フォーム・CAPTCHAなし"
    return "自動の候補", "自社フォーム・CAPTCHAなし"


def style(
    ws, widths: dict[str, int], wrap: set[str], status_col: str, status: list[str]
) -> None:
    head = PatternFill("solid", fgColor="1B3A2B")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(vertical="center", wrap_text=True)
    cols = [c.value for c in ws[1]]
    for i, name in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 12)
    wrap_letters = {get_column_letter(cols.index(n) + 1) for n in wrap if n in cols}
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(
                vertical="top", wrap_text=c.column_letter in wrap_letters
            )
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value
                c.font = Font(color="1155CC", underline="single")
    dv = DataValidation(
        type="list", formula1='"' + ",".join(status) + '"', allow_blank=True
    )
    ws.add_data_validation(dv)
    col = get_column_letter(cols.index(status_col) + 1)
    dv.add(f"{col}2:{col}{max(2, ws.max_row)}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions


def main(target_xlsx: str, enriched: str, gbiz: str, jassa: str, dst: str) -> None:
    ws = load_workbook(target_xlsx, read_only=True)["リスト"]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    targets = [
        d
        for d in (dict(zip(head, r)) for r in rows[1:])
        if d["区分"] in TARGET
        or (d["区分"] == "厚めHP×記事なし" and d["優先度"] == "高")
    ]
    contact = {r["domain"]: r for r in csv.DictReader(open(enriched, encoding="utf-8"))}
    info = {r["domain"]: r for r in csv.DictReader(open(gbiz, encoding="utf-8"))}
    jassa_addr = {
        dom(r["url"]): r["address"]
        for r in csv.DictReader(open(jassa, encoding="utf-8"))
    }
    # 従業員 300 人を超える会社は狙い（マーケ専任なし〜1人体制）から外れるので除く
    targets = [
        d
        for d in targets
        if not (
            info.get(d["ドメイン"], {}).get("従業員数", "").isdigit()
            and int(info[d["ドメイン"]]["従業員数"]) > 300
        )
    ]
    prio = {"高": 0, "中": 1, "低": 2}
    targets.sort(
        key=lambda d: (
            prio.get(d["優先度"], 3),
            TARGET.index(d["区分"]) if d["区分"] in TARGET else 3,
        )
    )

    wb = Workbook()
    call = wb.active
    call.title = "架電リスト"
    call_cols = [
        "No",
        "区分",
        "優先度",
        "会社名",
        "電話番号",
        "架電の一言目",
        "代表者",
        "従業員数",
        "所在地",
        "業種",
        "事業概要",
        "HP",
        "メディアURL",
        "記事数",
        "最終公開（推定）",
        "広告・MAタグ",
        "問い合わせURL",
        "フォームの送り方",
        "gBizINFO",
        "ステータス",
        "担当",
        "架電日",
        "結果メモ",
        "次のアクション",
    ]
    form = wb.create_sheet("フォーム送信リスト")
    form_cols = [
        "No",
        "区分",
        "優先度",
        "会社名",
        "送り方",
        "理由",
        "問い合わせURL",
        "フォームの種類",
        "入力欄",
        "必須の数",
        "用件の選択",
        "CAPTCHA",
        "確認画面",
        "同意チェック",
        "フォーム文面",
        "送信状態",
        "送信日",
        "メモ",
    ]
    call.append(call_cols)
    form.append(form_cols)
    for n, d in enumerate(targets, 1):
        dmn = d["ドメイン"]
        c = contact.get(dmn, {})
        g = info.get(dmn, {})
        name = g.get("正式社名") or d["サイト名"]
        hook = (d["架電の一言目"] or "").replace(f"{d['サイト名']}様", f"{name}様", 1)
        text = (d["フォーム文面"] or "").replace(d["サイト名"], name, 1)
        route, why = send_route(c)
        corp = g.get("法人番号", "")
        call.append(
            [
                n,
                d["区分"],
                d["優先度"],
                name,
                c.get("phone", ""),
                hook,
                g.get("代表者", ""),
                int(g["従業員数"]) if g.get("従業員数", "").isdigit() else None,
                g.get("所在地") or jassa_addr.get(dmn, ""),
                d["領域"],
                g.get("事業概要", ""),
                f"https://{dmn}",
                d["メディアURL"],
                d["記事数"],
                d["最終公開（推定）"],
                d["広告・MAタグ"],
                c.get("contact_url") or d["問い合わせURL"],
                route,
                f"https://info.gbiz.go.jp/hojin/ichiran?hojinBango={corp}"
                if corp
                else "",
                "未着手",
                "",
                "",
                "",
                "",
            ]
        )
        form.append(
            [
                n,
                d["区分"],
                d["優先度"],
                name,
                route,
                why,
                c.get("contact_url") or d["問い合わせURL"],
                c.get("form_kind", ""),
                c.get("form_fields", ""),
                int(c["required_count"])
                if c.get("required_count", "").isdigit()
                else None,
                "あり" if c.get("has_select") == "True" else "",
                c.get("captcha", ""),
                "あり" if c.get("confirm_step") == "True" else "",
                "あり" if c.get("privacy_consent") == "True" else "",
                (text + SIGNATURE)
                if route not in ("送らない", "フォームなし") and text
                else "",
                "送らない" if route == "送らない" else "未送信",
                "",
                "",
            ]
        )
    style(
        call,
        {
            "No": 5,
            "区分": 12,
            "優先度": 6,
            "会社名": 28,
            "電話番号": 14,
            "架電の一言目": 60,
            "代表者": 14,
            "所在地": 30,
            "事業概要": 40,
            "HP": 28,
            "メディアURL": 34,
            "問い合わせURL": 34,
            "フォームの送り方": 16,
            "gBizINFO": 20,
            "ステータス": 13,
            "結果メモ": 30,
            "次のアクション": 20,
        },
        {"架電の一言目", "事業概要", "結果メモ"},
        "ステータス",
        CALL_STATUS,
    )
    style(
        form,
        {
            "No": 5,
            "区分": 12,
            "優先度": 6,
            "会社名": 28,
            "送り方": 16,
            "理由": 26,
            "問い合わせURL": 36,
            "フォームの種類": 14,
            "入力欄": 30,
            "フォーム文面": 70,
            "送信状態": 12,
            "メモ": 24,
        },
        {"理由", "入力欄", "フォーム文面"},
        "送信状態",
        FORM_STATUS,
    )

    s = wb.create_sheet("集計と送り主")
    s.append(["項目", "件数"])
    s.append(["対象", len(targets)])
    s.append(
        [
            "電話番号あり",
            sum(1 for d in targets if contact.get(d["ドメイン"], {}).get("phone")),
        ]
    )
    routes: dict[str, int] = {}
    for d in targets:
        r, _ = send_route(contact.get(d["ドメイン"], {}))
        routes[r] = routes.get(r, 0) + 1
    for k, v in sorted(routes.items(), key=lambda x: -x[1]):
        s.append([f"送り方: {k}", v])
    s.append([])
    s.append(["フォームに入れる送り主の情報（記入してください）", ""])
    for k in (
        "会社名",
        "担当者名",
        "フリガナ",
        "部署・役職",
        "メール",
        "電話",
        "郵便番号",
        "住所",
        "URL",
    ):
        s.append([k, "株式会社MIRAIKU" if k == "会社名" else ""])
    s.append([])
    s.append(
        [
            "守ること",
            "営業お断りの会社には送らない／CAPTCHA は回避しない／不要と返信があれば exclude_domains.txt に足す／検索順位を約束しない",
        ]
    )
    s.column_dimensions["A"].width = 40
    s.column_dimensions["B"].width = 90
    wb.save(dst)
    print(dst, len(targets), routes)


if __name__ == "__main__":
    main(*sys.argv[1:6])
