"""問い合わせフォームの入力補助（半自動）。入力までを済ませ、送信ボタンは人が押す。

使い方:
  # 才木さんの PC で（ブラウザが開く。各社で内容を確かめて送信し、ターミナルで結果を入れる）
  python3 form_assist.py outreach.xlsx sender.json --limit 20

  # 入力の精度だけを確かめる（送信しない。画面を保存して、入れられなかった項目を出す）
  python3 form_assist.py outreach.xlsx sender.json --limit 5 --dry-run --shots shots/

このツールは送信ボタンを押さない。CAPTCHA には触れない（人が解く）。
「フォーム送信リスト」の送り方が「送らない」「フォームなし」「メールのみ」の会社と、送信状態が「未送信」以外の会社は飛ばす。
結果は outreach_log.csv（xlsx と同じフォルダ）に追記する。xlsx には書き込まない（開いたままでも動くように）。
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import Locator, Page, sync_playwright

SKIP_ROUTES = ("送らない", "フォームなし", "メールのみ")

# 各入力欄の手がかりを集める。own はその欄自身のラベル（ラジオの選択肢名など）、head はその欄の見出し
COLLECT_JS = r"""
(form) => {
  const els = [...form.querySelectorAll('input, textarea, select')];
  const txt = (n) => (n ? n.innerText || n.textContent || '' : '').replace(/\s+/g, ' ').trim();
  const ownOf = (el) => {
    if (el.id) { const l = form.querySelector(`label[for="${CSS.escape(el.id)}"]`); if (l) return txt(l); }
    const p = el.closest('label'); return p ? txt(p) : '';
  };
  // 見出しは、その欄（ラジオは同じ name の組）だけを囲む範囲から探す。ほかの欄まで囲んだら止める
  const headOf = (el) => {
    const mine = (x) => x === el || (el.type === 'radio' && x.type === 'radio' && x.name === el.name)
      || (el.type === 'checkbox' && x.type === 'checkbox' && x.name === el.name);
    let node = el;
    for (let k = 0; k < 6 && node && node !== form; k++) {
      node = node.parentElement;
      if (!node) break;
      if (node.matches('tr')) { const th = node.querySelector('th'); if (th) return txt(th); }
      if (node.matches('dd')) { const dt = node.previousElementSibling; if (dt) return txt(dt); }
      const others = [...node.querySelectorAll('input:not([type=hidden]), textarea, select')].filter(x => !mine(x));
      if (others.length) return '';
      for (const h of node.querySelectorAll('th, dt, legend, label, .label, .title, h3, h4, h5, p, span')) {
        if ([...node.querySelectorAll('input, textarea, select')].some(x => h.contains(x))) continue;
        const t = txt(h); if (t) return t;
      }
      const prev = node.previousElementSibling;
      if (prev && !prev.querySelector('input, textarea, select')) { const t = txt(prev); if (t && t.length < 40) return t; }
    }
    return '';
  };
  const shown = (el) => {
    const r = el.getBoundingClientRect(); const cs = getComputedStyle(el);
    return r.width > 2 && r.height > 2 && cs.visibility !== 'hidden' && cs.display !== 'none' && r.right > 0 && r.bottom > -5000;
  };
  return els.map((el, i) => {
    const own = ownOf(el).slice(0, 60), head = headOf(el).slice(0, 60);
    const box = el.closest('label') || el.parentElement;
    return {
      i, tag: el.tagName.toLowerCase(), type: (el.type || '').toLowerCase(), name: el.name || '', id: el.id || '',
      ph: el.placeholder || '', aria: el.getAttribute('aria-label') || '', own, head,
      // 自分のラベルがあればそれを優先し、見出しは補助にする（ラジオ・チェックボックスは組の見出しも使う）
      label: (el.type === 'radio' || el.type === 'checkbox' ? own + ' ' + head : (own || head)).trim(),
      req: el.required || el.getAttribute('aria-required') === 'true' || /必須|\*/.test(head + own),
      maxlen: el.maxLength || -1, ro: !!(el.readOnly || el.disabled),
      // チェックボックス・ラジオは見た目を作り込んで本体を隠すことが多いので、囲みの見え方で判断する
      visible: shown(el) || ((el.type === 'checkbox' || el.type === 'radio') && box && shown(box)),
      options: el.tagName === 'SELECT' ? [...el.options].map(o => o.text.trim()) : [],
      value: el.value || '',
    };
  });
}
"""

# いちばん入力欄の多いフォームの番号。form 要素が無いページ（作り込みのフォーム）は -2（body を使う）
FORM_PICK_JS = r"""
() => {
  const all = [...document.querySelectorAll('form')];
  const fields = (f) => f.querySelectorAll('input:not([type=hidden]):not([type=submit]):not([type=button]), select').length;
  const isSearch = (f) => f.getAttribute('role') === 'search' || /search/i.test(f.className + ' ' + f.id + ' ' + (f.action || ''))
    || !!f.querySelector('input[type=search], input[name=q], input[name=s], input[name=query]') && !f.querySelector('textarea');
  const forms = all.filter(f => !isSearch(f) && (f.querySelector('textarea') || fields(f) >= 4));
  forms.sort((a, b) => b.querySelectorAll('input, textarea, select').length - a.querySelectorAll('input, textarea, select').length);
  if (forms.length) return all.indexOf(forms[0]);
  return document.querySelector('textarea') && document.querySelectorAll('input').length >= 3 ? -2 : -1;
}
"""

# 求職者向けの応募・登録フォームの手がかり（営業の連絡を入れない）
APPLICANT = re.compile(
    r"自己PR|勤務希望|希望職種|希望勤務|登録希望|生年月日|年齢|履歴書|職務経歴"
)

CHOOSE = re.compile(
    r"その他|ご相談|ご提案|お問い?合わせ|サービス|営業|協業|パートナー|業務提携"
)
AVOID = re.compile(
    r"採用|求人|応募|エントリー|求職|登録|資料請求|見積|カタログ|サンプル|苦情|個人|情報を?(受け取|配信|お送り)|メルマガ|ニュースレター|に関する情報"
)


def kind_of(f: dict) -> str:
    key = " ".join([f["name"], f["id"], f["ph"], f["aria"], f["label"]])
    if f["type"] in ("hidden", "submit", "button", "image", "file", "reset", "search"):
        return ""
    if f["type"] == "checkbox":
        return (
            "consent"
            if re.search(
                r"同意|個人情報|プライバシー|privacy|agree|accept|policy",
                key,
                re.IGNORECASE,
            )
            else ""
        )
    if f["type"] == "radio" or f["tag"] == "select":
        group = " ".join([f["name"], f["head"]])
        if re.search(r"都道府県|pref", group, re.IGNORECASE):
            return "pref"
        if re.search(r"きっかけ|知った|経路|媒体|how", group, re.IGNORECASE):
            return "source"
        if re.search(r"役職|position", group, re.IGNORECASE):
            return "title"
        if re.search(
            r"連絡.{0,4}方法|ご連絡先の希望|返信方法|method|contact_?way",
            group,
            re.IGNORECASE,
        ):
            return "method"
        return "topic"
    if re.search(r"郵便|zip|postal|〒", key, re.IGNORECASE):
        return "zip"
    if re.search(r"都道府県|pref", key, re.IGNORECASE):
        return "pref"
    if f["type"] == "email" or (
        re.search(r"メール|e-?mail|mail", key, re.IGNORECASE)
        and not re.search(r"住所|所在地", key)
    ):
        return "email"
    if re.search(
        r"住所|所在地|address|addr|city|street|市区町村|番地", key, re.IGNORECASE
    ):
        return "address"
    if f["tag"] == "textarea":
        return "message"
    if f["type"] == "email" or re.search(r"メール|e-?mail|mail", key, re.IGNORECASE):
        return "email"
    kana = re.search(
        r"フリガナ|ふりがな|カナ|かな|kana|furigana|yomi", key, re.IGNORECASE
    )
    if kana and re.search(r"会社|社名|企業|法人", key):
        return "company_kana"
    if kana:
        if re.search(r"セイ|せい|姓|sei|last|family", key, re.IGNORECASE):
            return "last_kana"
        if re.search(
            r"メイ|めい|(?<![氏お])名(?!前)|mei|first|given", key, re.IGNORECASE
        ):
            return "first_kana"
        return "name_kana"
    if re.search(
        r"会社|社名|企業|法人|団体|組織|貴社|御社|company|organi[sz]ation|corp",
        key,
        re.IGNORECASE,
    ):
        return "company"
    if re.search(r"役職|position|title", key, re.IGNORECASE) and not re.search(
        r"件名|subject", key, re.IGNORECASE
    ):
        return "title"
    if re.search(r"部署|所属|department|division|post", key, re.IGNORECASE):
        return "department"
    if f["type"] == "tel" or re.search(r"電話|tel|phone|携帯", key, re.IGNORECASE):
        return "phone"
    if re.search(r"url|ホームページ|website|サイト", key, re.IGNORECASE):
        return "url"
    if re.search(r"件名|題名|subject|タイトル", key, re.IGNORECASE):
        return "subject"
    if re.search(
        r"(^|[^お氏])姓|last_?name|family|(^|[_\-\[])sei\b|name1|name_1|name_sei",
        key,
        re.IGNORECASE,
    ):
        return "last_name"
    if (
        re.search(
            r"^名$|first_?name|given|(^|[_\-\[])mei\b|name2|name_2|name_mei",
            key,
            re.IGNORECASE,
        )
        or f["label"].strip() == "名"
    ):
        return "first_name"
    if re.search(r"氏名|お名前|名前|ご担当|担当者|name", key, re.IGNORECASE):
        return "name"
    if re.search(
        r"内容|本文|ご質問|ご相談|message|inquiry|comment|body|詳細|ご用件",
        key,
        re.IGNORECASE,
    ):
        return "message"
    return ""


def pick_option(options: list[str], source: bool = False) -> int:
    cand = [
        (0 if re.search(r"その他", o) else 1 if CHOOSE.search(o) else 2, i)
        for i, o in enumerate(options)
        if o and not re.search(r"選択|選んで|^-+$|^—+$", o) and not AVOID.search(o)
    ]
    if source:
        cand = [(0 if re.search(r"その他", options[i]) else 1, i) for _, i in cand]
    cand.sort()
    return cand[0][1] if cand and cand[0][0] < 2 or (cand and source) else -1


def find_form(page: Page, wait_ms: int = 8000) -> tuple[Locator | None, str]:
    """入力欄のいちばん多いフォーム。外部フォームは遅れて出るので少し待ちながら探す"""
    for _ in range(max(1, wait_ms // 1000)):
        for fr in page.frames:
            try:
                idx = fr.evaluate(FORM_PICK_JS)
            except Exception:
                continue
            if idx >= 0:
                return fr.locator("form").nth(idx), ""
            if idx == -2:
                return fr.locator("body"), "form要素なし"
        page.wait_for_timeout(1000)
    return None, ""


def check_box(el: Locator) -> None:
    """チェックボックス・ラジオに印を付ける。

    マウスでは押さない。同意の文言の中のプライバシーポリシーのリンクを押してページを離れることがあるため、
    要素の click() を JS で呼び、効かなければ checked を直接立てて change を送る。
    """
    el.evaluate(
        "e => { if (!e.checked) e.click();"
        " if (!e.checked) { e.checked = true;"
        " e.dispatchEvent(new Event('input', {bubbles: true}));"
        " e.dispatchEvent(new Event('change', {bubbles: true})); } }",
        timeout=3000,
    )


def fill(
    page: Page, sender: dict, message: str, short: str = "", subject: str = ""
) -> tuple[list[str], list[str]]:
    """入力できた項目と、入力できなかった必須の項目を返す"""
    form, note = find_form(page)
    if form is None:
        return [], [
            "フォームが見つからない（外部フォームの読み込み待ちのことがある。画面で入力する）"
        ]
    fields = form.evaluate(COLLECT_JS)
    labels = " ".join(f["label"] + " " + f["name"] for f in fields)
    if APPLICANT.search(labels):
        return [], ["求職者向けの応募・登録フォームらしい（送らない）"]
    done, missing = [], ([note] if note else [])
    shown = [f for f in fields if f["visible"] and not f["ro"]]
    kinds = {f["i"]: kind_of(f) for f in shown}
    phone_parts = [f for f in shown if kinds[f["i"]] == "phone"]
    full = {
        "company": sender["company"],
        "company_kana": sender.get("company_kana", ""),
        "department": sender.get("department", ""),
        "title": sender.get("title", ""),
        "name": sender["last_name"] + " " + sender["first_name"],
        "last_name": sender["last_name"],
        "first_name": sender["first_name"],
        "name_kana": sender.get("last_kana", "") + " " + sender.get("first_kana", ""),
        "last_kana": sender.get("last_kana", ""),
        "first_kana": sender.get("first_kana", ""),
        "email": sender["email"],
        "url": sender.get("url", ""),
        "subject": subject or sender.get("subject", "オウンドメディア運用のご提案"),
    }
    filled: dict[str, int] = {}
    radios_done: set[str] = set()
    prev_zip = -1
    for f in shown:
        kind = kinds[f["i"]]
        el = form.locator("input, textarea, select").nth(f["i"])
        n = filled.get(kind, 0)
        try:
            if kind == "phone":
                if len(phone_parts) == 3:
                    el.fill(
                        sender["phone"].split("-")[phone_parts.index(f)], timeout=3000
                    )
                elif n == 0:
                    el.fill(
                        sender["phone"]
                        if f["type"] != "number"
                        and (f["maxlen"] in (-1, 0) or f["maxlen"] >= 12)
                        else sender["phone"].replace("-", ""),
                        timeout=3000,
                    )
                else:
                    continue
            elif kind == "zip":
                z = sender.get("zip", "")
                if n == 0:
                    el.fill(
                        z
                        if f["type"] != "number"
                        and (f["maxlen"] in (-1, 0) or f["maxlen"] >= 8)
                        else z.replace("-", ""),
                        timeout=3000,
                    )
                elif n == 1 and "-" in z and prev_zip >= 0:  # 3 桁と 4 桁に分かれた欄
                    form.locator("input, textarea, select").nth(prev_zip).fill(
                        z.split("-")[0], timeout=3000
                    )
                    el.fill(z.split("-")[1], timeout=3000)
                prev_zip = f["i"]
            elif kind == "pref":
                if f["tag"] == "select":
                    el.select_option(label=sender.get("pref", "東京都"), timeout=3000)
                elif f["type"] != "radio":
                    el.fill(sender.get("pref", ""), timeout=3000)
            elif kind == "address":
                if (
                    n == 0
                ):  # 住所が 2 欄（番地・建物名）に分かれていても、1 欄目にまとめて入れる
                    addr = sender.get("address", "")
                    if any(kinds[g["i"]] == "pref" for g in shown):
                        addr = addr.removeprefix(sender.get("pref", ""))
                    el.fill(addr, timeout=3000)
                else:
                    continue
            elif kind == "consent":
                check_box(el)
            elif kind == "message":
                if n > 0:
                    continue  # 本文は 1 か所だけ（ラベルの無い 2 つ目の欄は罠のことがある）
                text = short if short and 0 < f["maxlen"] < len(message) else message
                el.fill(text, timeout=5000)
                if text is short:
                    done.append("短い文面")
            elif kind == "title" and f["tag"] == "select":
                want = sender.get("title", "")
                i = next(
                    (k for k, o in enumerate(f["options"]) if want and want in o),
                    next((k for k, o in enumerate(f["options"]) if "その他" in o), -1),
                )
                if i >= 0:
                    el.select_option(index=i, timeout=3000)
                    done.append(f"役職:{f['options'][i][:10]}")
                elif f["req"]:
                    missing.append("役職の選択肢")
                continue
            elif kind in ("topic", "source", "method") and f["tag"] == "select":
                if kind == "method":
                    i = next(
                        (k for k, o in enumerate(f["options"]) if "メール" in o), -1
                    )
                else:
                    i = pick_option(f["options"], source=kind == "source")
                if i >= 0:
                    el.select_option(index=i, timeout=3000)
                    done.append(f"選択:{f['options'][i][:20]}")
                elif f["req"]:
                    missing.append(f"選択肢（{f['head'][:20]}）")
                continue
            elif kind in ("topic", "source", "method") and f["type"] == "radio":
                if f["name"] in radios_done:
                    continue
                radios_done.add(f["name"])
                group = [
                    g for g in shown if g["type"] == "radio" and g["name"] == f["name"]
                ]
                opts = [(g["own"] or g["value"]) for g in group]
                if kind == "method":
                    i = next((k for k, o in enumerate(opts) if "メール" in o), -1)
                else:
                    i = pick_option(opts, source=kind == "source")
                if i >= 0:
                    check_box(
                        form.locator("input, textarea, select").nth(group[i]["i"])
                    )
                    done.append(f"選択:{opts[i][:20]}")
                elif f["req"]:
                    missing.append(f"選択肢（{f['head'][:20]}）")
                continue
            elif kind in full:
                if full[kind] and (n == 0 or kind == "email"):
                    el.fill(full[kind], timeout=3000)
                elif not full[kind] and f["req"]:
                    missing.append(f"{kind}（送り主情報が空）")
                    continue
                else:
                    continue
            else:
                if f["req"] and f["type"] not in ("checkbox", "radio"):
                    missing.append(
                        f"不明な必須項目（{(f['label'] or f['name'])[:30]}）"
                    )
                continue
            filled[kind] = n + 1
            if n == 0:
                done.append(kind)
        except Exception as e:  # 1 項目の失敗で止めない
            missing.append(f"{kind or f['name']}: {type(e).__name__}")
    return done, missing


def rows_from(xlsx: str) -> list[dict]:
    ws = load_workbook(xlsx, read_only=True)["フォーム送信リスト"]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:]]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("sender")
    ap.add_argument("--limit", type=int, default=20)
    ap.add_argument(
        "--dry-run",
        action="store_true",
        help="入力だけ試して画面を保存する。送信しない",
    )
    ap.add_argument("--shots", default="shots")
    ap.add_argument("--start", type=int, default=1, help="この No から始める")
    ap.add_argument("--executable", default=None, help="Chromium の場所（任意）")
    a = ap.parse_args()

    sender = json.loads(Path(a.sender).read_text(encoding="utf-8"))
    log_path = Path(a.xlsx).with_name("outreach_log.csv")
    logged = set()
    if log_path.exists():
        with open(log_path, encoding="utf-8") as fh:
            logged = {
                r["No"]
                for r in csv.DictReader(fh)
                if r["結果"] in ("送信した", "送らない")
            }
    todo = [
        r
        for r in rows_from(a.xlsx)
        if r["送り方"] not in SKIP_ROUTES
        and r["送信状態"] == "未送信"
        and r["フォーム文面"]
        and str(r["No"]) not in logged
        and int(r["No"]) >= a.start
    ][: a.limit]
    Path(a.shots).mkdir(exist_ok=True)
    new_log = not log_path.exists()
    with (
        sync_playwright() as p,
        open(log_path, "a", newline="", encoding="utf-8") as lf,
    ):
        log = csv.writer(lf)
        if new_log:
            log.writerow(
                [
                    "日時",
                    "No",
                    "会社名",
                    "問い合わせURL",
                    "結果",
                    "入力した項目",
                    "入力できなかった項目",
                    "メモ",
                ]
            )
        browser = p.chromium.launch(headless=a.dry_run, executable_path=a.executable)
        ctx = browser.new_context(
            locale="ja-JP", viewport={"width": 1280, "height": 900}
        )
        for r in todo:
            page = ctx.new_page()
            print(f"\n[{r['No']}] {r['会社名']}  {r['問い合わせURL']}")
            try:
                page.goto(
                    r["問い合わせURL"], wait_until="domcontentloaded", timeout=30000
                )
                page.wait_for_timeout(2000)
                done, missing = fill(
                    page,
                    sender,
                    r["フォーム文面"],
                    r.get("短い文面") or "",
                    r.get("件名") or "",
                )
            except Exception as e:
                done, missing = [], [f"開けない: {type(e).__name__}"]
            print("  入力:", ", ".join(done) or "なし")
            if missing:
                print("  要確認:", " / ".join(missing))
            if "message" not in done and not a.dry_run:
                # 本文を入れられなかったときは、手で貼れるように出しておく
                print("\n----- 件名 -----\n" + (r.get("件名") or ""))
                print("----- 文面 -----\n" + r["フォーム文面"] + "----------------")
            if a.dry_run:
                page.screenshot(
                    path=str(
                        Path(a.shots) / f"{r['No']:04d}.png"
                        if isinstance(r["No"], int)
                        else f"{a.shots}/{r['No']}.png"
                    ),
                    full_page=True,
                )
                log.writerow(
                    [
                        datetime.now().isoformat(timespec="seconds"),
                        r["No"],
                        r["会社名"],
                        r["問い合わせURL"],
                        "試験入力",
                        ",".join(done),
                        " / ".join(missing),
                        "",
                    ]
                )
                lf.flush()
                page.close()
                continue
            print(
                "  → ブラウザで内容を確かめ、必要なら直し、CAPTCHA があれば解いて、送信ボタンを押してください。"
            )
            ans = (
                input("  送信した=Enter / 飛ばす=s / 送らない=n / 終了=q : ")
                .strip()
                .lower()
            )
            result = {
                "": "送信した",
                "s": "飛ばした",
                "n": "送らない",
                "q": "飛ばした",
            }.get(ans, "飛ばした")
            memo = input("  メモ（任意）: ").strip() if ans != "q" else ""
            log.writerow(
                [
                    datetime.now().isoformat(timespec="seconds"),
                    r["No"],
                    r["会社名"],
                    r["問い合わせURL"],
                    result,
                    ",".join(done),
                    " / ".join(missing),
                    memo,
                ]
            )
            lf.flush()
            page.close()
            if ans == "q":
                break
        browser.close()
    print(f"\n記録: {log_path}")


if __name__ == "__main__":
    main()
